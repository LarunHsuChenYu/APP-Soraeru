# -*- coding: utf-8 -*-
"""Generate image-first Traditional Chinese user manual DOCX for 空耳聯合國."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parent
HTML_DIR = ROOT.parent / "AI 空耳外語學習 APP－MVP 系統規劃書" / "stitch_soraeru_mnemonic_vocabulary_app" / "html"
LOGO = ROOT.parent / "AI 空耳外語學習 APP－MVP 系統規劃書" / "stitch_soraeru_mnemonic_vocabulary_app" / "png" / "logo.png"
SHOT_DIR = ROOT / "screenshots"
OUT_DOCX = ROOT / "空耳聯合國_簡易使用說明書.docx"

# Mobile phone frame size used by Stitch prototypes
VIEWPORT = {"width": 430, "height": 900}

SCREENS = [
    ("L00_code.html", "l00_splash.png"),
    ("L01_code.html", "l01_login.png"),
    ("L02_code.html", "l02_register.png"),
    ("l03_forgot_password_screen.html", "l03_forgot.png"),
    ("l04_onboarding_screen_mvp_rev.html", "l04_onboarding.png"),
    ("l05_home_screen_mvp_rev.html", "l05_home.png"),
    ("l06_word_input_screen.html", "l06_input.png"),
    ("l07_image_pick_screen.html", "l07_image.png"),
    ("l08_ocr_select_screen.html", "l08_ocr.png"),
    ("l09_analyzing_screen.html", "l09_analyzing.png"),
    ("l10_analysis_result_mvp_rev.html", "l10_result.png"),
    ("l11_notebook_list_mvp_rev.html", "l11_notebook.png"),
    ("l12_notebook_detail_screen.html", "l12_detail.png"),
    ("l13_settings_screen_mvp_rev.html", "l13_settings.png"),
]


def capture_screens() -> None:
    SHOT_DIR.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as p:
        browser = p.chromium.launch(channel="chrome", headless=True)
        context = browser.new_context(
            viewport=VIEWPORT,
            device_scale_factor=2,
            locale="zh-TW",
        )
        page = context.new_page()
        for html_name, png_name in SCREENS:
            html_path = HTML_DIR / html_name
            if not html_path.exists():
                raise FileNotFoundError(html_path)
            url = html_path.as_uri()
            page.goto(url, wait_until="load", timeout=90_000)
            page.wait_for_timeout(1500)
            out = SHOT_DIR / png_name
            page.screenshot(path=str(out), full_page=True)
            print(f"captured {png_name}")
        browser.close()


def set_run_font(run, name: str = "Microsoft JhengHei", size: int | None = None, bold: bool = False, color: RGBColor | None = None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_heading_zh(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=18 if level == 1 else 14, bold=True, color=RGBColor(0x00, 0x4D, 0x64))


def add_body(doc: Document, text: str, *, center: bool = False, size: int = 11, bold: bool = False, space_after: int = 6) -> None:
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    set_run_font(run, size=10, color=RGBColor(0x3F, 0x48, 0x4D))


def add_phone_image(doc: Document, png: Path, width_cm: float = 7.2) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(png), width=Cm(width_cm))


def add_step_block(doc: Document, title: str, steps: list[str], image: Path, tip: str | None = None) -> None:
    add_heading_zh(doc, title, level=2)
    for i, step in enumerate(steps, 1):
        add_body(doc, f"{i}. {step}", size=11, space_after=2)
    if tip:
        add_body(doc, f"提示：{tip}", size=10, space_after=6)
    add_phone_image(doc, image)
    add_caption(doc, "對應畫面示意（設計稿）")


def build_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)

    # Cover
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(40)
        run = p.add_run()
        run.add_picture(str(LOGO), width=Cm(4.5))

    add_body(doc, "空耳聯合國", center=True, size=28, bold=True, space_after=4)
    add_body(doc, "簡易使用說明書", center=True, size=18, bold=True, space_after=8)
    add_body(doc, "用發音，記住外語", center=True, size=12, space_after=4)
    add_body(doc, "版本 1.0｜Android｜繁體中文", center=True, size=10, space_after=18)
    add_body(
        doc,
        "本說明以畫面為主。請對照手機上的實際按鈕操作；圖片為介面示意。",
        center=True,
        size=10,
        space_after=20,
    )

    # What is it
    add_heading_zh(doc, "這是什麼 App？")
    add_body(
        doc,
        "輸入或拍照取得外語單字，用 AI 產生華語「空耳」近似音幫助記憶。"
        "正式發音請以 App 提供的讀音／播放為準；近似音只是助記用。",
        size=11,
    )
    add_phone_image(doc, SHOT_DIR / "l00_splash.png", width_cm=6.5)
    add_caption(doc, "啟動畫面")

    # Bottom tabs overview
    add_heading_zh(doc, "主畫面導覽")
    add_body(doc, "登入後底部有三個分頁：", size=11)
    add_body(doc, "• 首頁：輸入單字、拍照取字、查看今日額度", size=11, space_after=2)
    add_body(doc, "• 單字本：收藏的單字卡列表", size=11, space_after=2)
    add_body(doc, "• 設定：帳號、額度、標記偏好、教學與法律文件", size=11, space_after=6)
    add_phone_image(doc, SHOT_DIR / "l05_home.png")
    add_caption(doc, "首頁（底部可切換「首頁／單字本／設定」）")

    # Login
    add_step_block(
        doc,
        "一、登入或註冊",
        [
            "開啟 App，進入登入頁。",
            "可用 Email／密碼登入，或使用 Google 登入（Android）。",
            "沒有帳號請點註冊；忘記密碼可走重設流程。",
            "首次登入會看到三步使用說明，看完後點「開始使用」。",
        ],
        SHOT_DIR / "l01_login.png",
        tip="登入後才可進行 AI 分析與寫入單字卡。",
    )
    add_phone_image(doc, SHOT_DIR / "l04_onboarding.png", width_cm=6.8)
    add_caption(doc, "首次使用說明（之後可在設定重新觀看）")

    # Manual input flow
    add_step_block(
        doc,
        "二、手動輸入單字並分析",
        [
            "在首頁點「輸入單字」。",
            "輸入外語單字或短語，來源語言可選「自動偵測」。",
            "點「開始分析」，等待語言偵測、詞義與近似音產生。",
            "在結果頁聆聽正式發音，挑選合適的空耳近似音。",
            "點「儲存單字卡」收藏到單字本。",
        ],
        SHOT_DIR / "l06_input.png",
        tip="若該字已收藏過，可能直接開啟既有單字卡，不另扣額度。",
    )
    add_phone_image(doc, SHOT_DIR / "l09_analyzing.png", width_cm=6.8)
    add_caption(doc, "分析進行中")
    add_phone_image(doc, SHOT_DIR / "l10_result.png", width_cm=6.8)
    add_caption(doc, "分析結果：正式讀音＋空耳候選（請留意「僅供記憶」提示）")

    # OCR flow
    add_step_block(
        doc,
        "三、拍照／選圖取字",
        [
            "在首頁點「拍照／選擇圖片」。",
            "用相機拍攝或從相簿選圖（可依需要選擇語系別）。",
            "點開始辨識：文字辨識在手機本機完成，原圖不上傳。",
            "在選字畫面點選或編輯要分析的單字。",
            "之後流程與手動輸入相同：分析 → 選空耳 → 儲存。",
        ],
        SHOT_DIR / "l07_image.png",
        tip="同一張圖可稍後在單字卡詳情繼續選其他字。",
    )
    add_phone_image(doc, SHOT_DIR / "l08_ocr.png", width_cm=6.8)
    add_caption(doc, "辨識後選擇要分析的單字")

    # Notebook
    add_step_block(
        doc,
        "四、單字本與單字卡",
        [
            "底部切換到「單字本」，或從首頁進入「我的單字卡」。",
            "可用搜尋或語言篩選找卡片。",
            "點開單字卡可播放正式發音、查看／編輯我的近似音。",
            "需要時可「重新分析」或「刪除單字卡」。",
        ],
        SHOT_DIR / "l11_notebook.png",
    )
    add_phone_image(doc, SHOT_DIR / "l12_detail.png", width_cm=6.8)
    add_caption(doc, "單字卡詳情")

    # Settings
    add_step_block(
        doc,
        "五、設定與額度",
        [
            "打開「設定」查看帳號與「今日剩餘 AI 次數」。",
            "標記偏好可選注音／拼音／混合，套用到新產生的卡片。",
            "可重新觀看教學、閱讀隱私權政策與 AI 內容聲明。",
            "登出或刪除帳號請在設定頁操作。",
        ],
        SHOT_DIR / "l13_settings.png",
        tip="每次成功的新分析會消耗當日 AI 額度。",
    )

    # Quick tips
    add_heading_zh(doc, "使用小提醒")
    tips = [
        "空耳近似音＝助記諧音，不是正式音標。",
        "結果頁若標示「AI 草稿」，表示尚未經聽感核定，請以正式發音為準。",
        "拍照取字在本機辨識，原圖不會上傳伺服器。",
        "免費方案有每日 AI 次數上限，用完請隔日再試。",
        "本說明畫面取自設計示意，實際版面可能因版本略有差異。",
    ]
    for t in tips:
        add_body(doc, f"• {t}", size=11, space_after=3)

    add_heading_zh(doc, "附錄：其他相關畫面")
    add_body(doc, "註冊", size=12, bold=True, space_after=2)
    add_phone_image(doc, SHOT_DIR / "l02_register.png", width_cm=6.5)
    add_caption(doc, "建立帳號")
    add_body(doc, "忘記密碼", size=12, bold=True, space_after=2)
    add_phone_image(doc, SHOT_DIR / "l03_forgot.png", width_cm=6.5)
    add_caption(doc, "重設密碼")

    add_body(doc, "— 空耳聯合國 簡易使用說明書 完 —", center=True, size=10, space_after=0)

    doc.save(str(OUT_DOCX))
    print(f"wrote {OUT_DOCX}")


def main() -> None:
    capture_screens()
    build_docx()


if __name__ == "__main__":
    main()
